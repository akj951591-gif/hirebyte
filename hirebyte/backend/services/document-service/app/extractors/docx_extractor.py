from __future__ import annotations

import io
import zipfile

import docx
from docx.document import Document

from . import ExtractedText, ExtractionError, clean_text


def _table_text(document: Document) -> list[str]:
    rows: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
    return rows


def extract_docx(data: bytes) -> ExtractedText:
    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        # Password-protected Word files are encrypted containers, not ZIPs, so they land here too.
        raise ExtractionError("CORRUPT_FILE", "This DOCX file is damaged, password-protected or not a real Word document.") from exc

    paragraphs = [p.text for p in document.paragraphs]
    text, truncated = clean_text("\n".join(paragraphs + _table_text(document)))
    if not text:
        raise ExtractionError("NO_TEXT", "This document has no text.")
    warnings = ["The document was long, so only the beginning was kept."] if truncated else []
    # Word files have no fixed pages; report 1 so callers have a consistent number.
    return ExtractedText(text=text, page_count=1, truncated=truncated, warnings=warnings)
