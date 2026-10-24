import io

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def make_pdf(*lines: str, password: str | None = None) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    for i, line in enumerate(lines):
        page.insert_text((72, 72 + 18 * i), line)
    if password:
        return doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password)
    return doc.tobytes()


def make_docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def post(data: bytes, filename: str, **headers: str):
    return client.post("/extract", files={"file": (filename, data, "application/octet-stream")}, headers=headers)


def test_health():
    assert client.get("/health").json() == {"status": "ok", "ocr": False}


def test_extracts_pdf_text():
    res = post(make_pdf("Jane Doe", "Skills:   Python,  SQL"), "resume.pdf")
    assert res.status_code == 200
    body = res.json()
    assert body["text"] == "Jane Doe\nSkills: Python, SQL"
    assert body["pageCount"] == 1
    assert body["charCount"] == len(body["text"])
    assert body["truncated"] is False


def test_extracts_docx_paragraphs_and_tables():
    res = post(make_docx("Jane Doe", "Backend engineer", table=[["Skill", "Years"], ["Go", "3"]]), "resume.docx")
    assert res.status_code == 200
    assert res.json()["text"] == "Jane Doe\nBackend engineer\nSkill | Years\nGo | 3"


def test_detects_type_by_content_not_extension():
    # A PDF renamed to .docx is still read as a PDF; a text file renamed to .pdf is refused.
    assert post(make_pdf("Real PDF text here"), "renamed.docx").json()["text"] == "Real PDF text here"
    res = post(b"just some text", "fake.pdf")
    assert res.status_code == 415
    assert res.json()["error"]["code"] == "UNSUPPORTED_TYPE"


@pytest.mark.parametrize(
    ("data", "filename", "status", "code"),
    [
        (b"", "empty.pdf", 400, "EMPTY_FILE"),
        (b"%PDF-1.7 broken", "broken.pdf", 422, "CORRUPT_FILE"),
        (b"PK\x03\x04not-a-real-zip", "broken.docx", 422, "CORRUPT_FILE"),
        (make_pdf(), "scanned.pdf", 422, "NO_TEXT"),
        (make_docx(), "empty.docx", 422, "NO_TEXT"),
    ],
    ids=["empty", "corrupt-pdf", "corrupt-docx", "scanned-pdf", "empty-docx"],
)
def test_refuses_unreadable_documents(data, filename, status, code):
    res = post(data, filename)
    assert res.status_code == status
    assert res.json()["error"]["code"] == code


def test_refuses_password_protected_pdf():
    res = post(make_pdf("secret resume", password="pw"), "locked.pdf")
    assert res.status_code == 422
    assert res.json()["error"]["code"] == "PASSWORD_PROTECTED"


def test_refuses_files_over_the_limit(monkeypatch):
    monkeypatch.setattr("app.routes.extract.MAX_BYTES", 1000)
    res = post(make_pdf("x" * 50) + b"\0" * 2000, "big.pdf")
    assert res.status_code == 413
    assert res.json()["error"]["code"] == "FILE_TOO_LARGE"


def test_truncates_very_long_documents(monkeypatch):
    monkeypatch.setattr("app.extractors.MAX_CHARS", 10)
    body = post(make_pdf("abcdefghijklmnopqrstuvwxyz"), "long.pdf").json()
    assert body["text"] == "abcdefghij"
    assert body["truncated"] is True
    assert body["warnings"]


def test_requires_the_service_token_when_configured(monkeypatch):
    monkeypatch.setenv("DOCUMENT_SERVICE_TOKEN", "shared-secret")
    assert post(make_pdf("hello there"), "a.pdf").status_code == 401
    assert post(make_pdf("hello there"), "a.pdf", **{"X-Service-Token": "wrong"}).status_code == 401
    assert post(make_pdf("hello there"), "a.pdf", **{"X-Service-Token": "shared-secret"}).status_code == 200


def test_missing_file_field():
    res = client.post("/extract", data={"nope": "1"})
    assert res.status_code == 400
    assert res.json()["error"]["code"] == "VALIDATION_ERROR"
