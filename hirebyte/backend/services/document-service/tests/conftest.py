import io

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app


@pytest.fixture()
def client():
    return TestClient(app)


def make_docx(paragraphs=(), table_rows=()) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf(pages=(), encrypt_with=None) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
        else:  # "scanned-like": a page with graphics but no text layer
            page.draw_rect(pymupdf.Rect(50, 50, 200, 200))
    if encrypt_with:
        return doc.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256,
            user_pw=encrypt_with,
            owner_pw=encrypt_with + "-owner",
        )
    return doc.tobytes()


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"